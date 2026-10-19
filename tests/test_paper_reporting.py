"""Tests for paper-only Telegram wording and Excel report export."""
import csv
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from app import notifications, storage


class PaperReportingTests(unittest.TestCase):
    def test_education_header_on_entry_and_close_alerts(self):
        trade = {
            "trade_id": "TEST-001", "symbol": "TEST", "side": "LONG",
            "signal_time": "2026-10-20T10:00:00+05:30",
            "entry_price": 100, "stop_loss": 98, "target": 103,
            "quantity": 10, "notional": 1000, "nifty_trend": "POSITIVE",
            "stock_poc": 99, "setup": "Sweep>Displacement>MSS>Retest",
        }
        entry = notifications.entry_alert(trade)
        close = notifications.close_alert(
            {**trade, "result": "WIN", "exit_reason": "TP", "exit_price": 103, "net_pnl": 30},
            {"closed": 1, "win_rate_pct": 100, "gross_profit": 30, "gross_loss": 0, "profit_factor": float("inf")},
        )
        self.assertIn("EDUCATION PURPOSE ONLY", entry)
        self.assertIn("EDUCATION PURPOSE ONLY", close)
        self.assertIn("PAPER TRADING ONLY", entry)
        self.assertIn("NO REAL ORDERS", close)

    def test_excel_export_contains_required_sheets(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            trade_path = root / "trade_history.csv"
            log_path = root / "system_log.csv"
            state_path = root / "paper_state.json"
            with patch.object(storage, "DATA_DIR", root), \
                 patch.object(storage, "TRADE_FILE", trade_path), \
                 patch.object(storage, "LOG_FILE", log_path), \
                 patch.object(storage, "STATE_FILE", state_path):
                trade = {
                    "trade_id": "TEST-001", "symbol": "TEST", "side": "LONG",
                    "signal_time": "2026-10-20T10:00:00+05:30",
                    "entry_time": "2026-10-20T10:15:00+05:30", "entry_price": 100,
                    "quantity": 10, "notional": 1000, "stop_loss": 98, "target": 103,
                    "exit_time": "2026-10-20T10:45:00+05:30", "exit_price": 103,
                    "exit_reason": "TP", "gross_pnl": 30, "result": "WIN",
                    "nifty_trend": "POSITIVE", "stock_poc": 99, "setup": "Sweep>Retest",
                    "estimated_cost": 0, "net_pnl": 30,
                }
                storage.save_trade(trade)
                storage.log_event("INFO", "TEST", "report test")
                path = storage.export_excel()
                wb = load_workbook(path, read_only=True)
                self.assertEqual(
                    set(wb.sheetnames),
                    {"TRADE_HISTORY", "CONDITIONS", "DAILY_SUMMARY", "PERFORMANCE", "SYSTEM_LOG"},
                )
                self.assertEqual(wb["TRADE_HISTORY"].max_row, 2)
                self.assertEqual(wb["DAILY_SUMMARY"].max_row, 2)
                self.assertEqual(wb["PERFORMANCE"].cell(2, 1).value, 1)
                wb.close()


if __name__ == "__main__":
    unittest.main()
