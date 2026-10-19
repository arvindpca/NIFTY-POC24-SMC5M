"""Restart-safe storage for paper-only positions and audit history."""
import csv
import json
import os
from pathlib import Path

DATA_DIR = Path(os.getenv("DATA_DIR", "data"))
DATA_DIR.mkdir(parents=True, exist_ok=True)
STATE_FILE = DATA_DIR / "paper_state.json"
TRADE_FILE = DATA_DIR / "trade_history.csv"
LOG_FILE = DATA_DIR / "system_log.csv"
FIELDS = [
    "trade_id", "symbol", "side", "signal_time", "entry_time", "entry_price",
    "quantity", "notional", "stop_loss", "target", "exit_time", "exit_price",
    "exit_reason", "gross_pnl", "result", "nifty_trend", "stock_poc",
    "setup", "estimated_cost", "net_pnl"
]


def load_state():
    if not STATE_FILE.exists():
        return {"open_positions": {}, "seen_signals": [], "last_candle": {}}
    with STATE_FILE.open(encoding="utf-8") as f:
        state = json.load(f)
    state.setdefault("open_positions", {})
    state.setdefault("seen_signals", [])
    state.setdefault("last_candle", {})
    return state


def save_state(state):
    tmp = STATE_FILE.with_suffix(".tmp")
    with tmp.open("w", encoding="utf-8") as f:
        json.dump(state, f, indent=2, ensure_ascii=False)
    os.replace(tmp, STATE_FILE)


def append_csv(path, fields, record):
    new_file = not path.exists()
    with path.open("a", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        if new_file:
            writer.writeheader()
        writer.writerow(record)


def save_trade(trade):
    """Append each closed virtual trade once, even after an interrupted workflow run."""
    trade_id = str(trade.get("trade_id", "")).strip()
    if trade_id and TRADE_FILE.exists():
        with TRADE_FILE.open(newline="", encoding="utf-8-sig") as f:
            for row in csv.DictReader(f):
                if row.get("trade_id") == trade_id:
                    return False
    append_csv(TRADE_FILE, FIELDS, trade)
    return True


def log_event(level, event, detail=""):
    append_csv(LOG_FILE, ["timestamp", "level", "event", "detail"], {
        "timestamp": __import__("datetime").datetime.now().astimezone().isoformat(),
        "level": level, "event": event, "detail": str(detail)[:1000]
    })



def export_excel():
    """Create a readable Excel workbook from the persisted CSV audit trail."""
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    from datetime import datetime

    DATA_DIR.mkdir(parents=True, exist_ok=True)
    out_path = DATA_DIR / "strategy_b_paper_history.xlsx"
    wb = Workbook()
    default = wb.active
    wb.remove(default)

    def read_rows(path):
        if not path.exists():
            return [], []
        with path.open(newline="", encoding="utf-8-sig") as f:
            reader = csv.DictReader(f)
            rows = list(reader)
            return reader.fieldnames or [], rows

    trade_headers, trades = read_rows(TRADE_FILE)
    log_headers, logs = read_rows(LOG_FILE)

    def sheet_from_rows(title, headers, rows):
        ws = wb.create_sheet(title)
        if headers:
            ws.append(list(headers))
            for row in rows:
                ws.append([row.get(h, "") for h in headers])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="244062")
        for col in ws.columns:
            values = list(col)
            width = min(max(max((len(str(c.value or "")) for c in values), default=10) + 2, 10), 42)
            ws.column_dimensions[get_column_letter(col[0].column)].width = width
        return ws

    sheet_from_rows("TRADE_HISTORY", trade_headers, trades)
    condition_fields = [f for f in ("trade_id", "symbol", "side", "signal_time", "nifty_trend", "stock_poc", "setup") if f in trade_headers]
    sheet_from_rows("CONDITIONS", condition_fields, trades)

    pnl_values = []
    for trade in trades:
        try:
            pnl_values.append((trade, float(trade.get("net_pnl") or 0)))
        except (TypeError, ValueError):
            continue
    daily = {}
    for trade, pnl in pnl_values:
        day = (trade.get("exit_time") or trade.get("entry_time") or "Unknown")[:10]
        item = daily.setdefault(day, {"trades": 0, "wins": 0, "losses": 0, "breakeven": 0, "gross_profit": 0.0, "gross_loss": 0.0, "net_pnl": 0.0})
        item["trades"] += 1
        item["wins"] += int(pnl > 0)
        item["losses"] += int(pnl < 0)
        item["breakeven"] += int(pnl == 0)
        item["gross_profit"] += max(pnl, 0)
        item["gross_loss"] += max(-pnl, 0)
        item["net_pnl"] += pnl
    daily_headers = ["date", "closed_trades", "wins", "losses", "breakeven", "win_rate_pct", "gross_profit", "gross_loss", "net_pnl"]
    daily_rows = []
    for day, item in sorted(daily.items()):
        daily_rows.append({
            "date": day, **{k: item[k] for k in ("trades", "wins", "losses", "breakeven", "gross_profit", "gross_loss", "net_pnl")},
            "closed_trades": item["trades"],
            "win_rate_pct": round(100 * item["wins"] / item["trades"], 2) if item["trades"] else 0,
        })
    sheet_from_rows("DAILY_SUMMARY", daily_headers, daily_rows)

    wins = sum(max(pnl, 0) for _, pnl in pnl_values)
    losses = sum(max(-pnl, 0) for _, pnl in pnl_values)
    performance = [{
        "closed_trades": len(pnl_values),
        "wins": sum(pnl > 0 for _, pnl in pnl_values),
        "losses": sum(pnl < 0 for _, pnl in pnl_values),
        "breakeven": sum(pnl == 0 for _, pnl in pnl_values),
        "win_rate_pct": round(100 * sum(pnl > 0 for _, pnl in pnl_values) / len(pnl_values), 2) if pnl_values else 0,
        "gross_profit": round(wins, 2),
        "gross_loss": round(losses, 2),
        "profit_factor": round(wins / losses, 4) if losses else ("INF" if wins else 0),
        "net_pnl": round(sum(pnl for _, pnl in pnl_values), 2),
        "generated_at_ist": datetime.now().astimezone().isoformat(),
    }]
    sheet_from_rows("PERFORMANCE", list(performance[0].keys()), performance)
    sheet_from_rows("SYSTEM_LOG", log_headers, logs)
    wb.save(out_path)
    return out_path
